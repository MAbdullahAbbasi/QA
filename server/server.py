from flask import *
import os
import tempfile
from openpyxl import Workbook
from openpyxl.styles import PatternFill
import csv
import uuid
import assemblyai as aai
import google.generativeai as genai
import requests
from io import StringIO, BytesIO
from werkzeug.utils import secure_filename

# Setup Flask app
template_dir = os.path.abspath(os.path.join(os.path.dirname(__file__), 'templates'))
app = Flask(__name__, template_folder=template_dir)

EXPORT_FOLDER = os.path.join(os.getcwd(), 'exports')
os.makedirs(EXPORT_FOLDER, exist_ok=True)

UPLOAD_FOLDER = os.path.join(os.getcwd(), 'uploads')
os.makedirs(UPLOAD_FOLDER, exist_ok=True)

# API Keys
aai.settings.api_key = "f2b25c07671e4d2cac25479469debcab"
genai.configure(api_key="AIzaSyC3akknXGbeuvAg_kBRkGsi586RuXeHkHo")

@app.route("/")
def home():
    return render_template("main.html")

@app.route('/process_folder', methods=['POST'])
def process_folder():
    uploaded_files = request.files.getlist('audios')
    if not uploaded_files:
        return jsonify({"error": "No files uploaded"}), 400

    questions = [
        "What is the year make and model of your vehicle?",
        "How many miles does your vehicle have",
        "Can you spell your first and last name, please",
        "In witch state do you currently reside",
        "May I have your email address, please"
    ]

    headers = [
        "File Name",
        "Vehicle Info",
        "Mileage",
        "Full Name",
        "State",
        "Email",
        "Qualification Status"
    ]

    results = []

    for audio_file in uploaded_files:
        filename = secure_filename(audio_file.filename)

        with tempfile.NamedTemporaryFile(delete=False, suffix='.wav') as temp_audio:
            audio_file.save(temp_audio.name)
            transcribed_text = audio_transcription(temp_audio.name)
        os.remove(temp_audio.name)

        prompt = build_prompt(transcribed_text, questions)
        filtered_answers = resource_generation_gemini(prompt)
        qualification = is_qualified_gemini(filtered_answers)

        # Extract answers in order of questions
        answers = [filtered_answers.get(q.lower(), '') for q in questions]
        row = [filename] + [str(ans) if not isinstance(ans, str) else ans for ans in answers] + [str(qualification)]
        # row = [filename] + answers + [qualification]
        results.append(row)

    # Generate Excel file saved on disk with unique name
    filename = f'results_{uuid.uuid4().hex}.xlsx'
    excel_path = os.path.join(EXPORT_FOLDER, filename)
    save_excel(headers, results, excel_path)

    # Return JSON: data + download link to Excel file
    return jsonify({
        "headers": headers,
        "rows": results,
        "excel_download_url": url_for('download_excel', filename=filename)
    })

@app.route('/download_excel/<filename>')
def download_excel(filename):
    file_path = os.path.join(EXPORT_FOLDER, filename)
    if not os.path.exists(file_path):
        return jsonify({"error": "File not found"}), 404
    return send_file(file_path, mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
                     as_attachment=True, download_name='results.xlsx')

def save_excel(headers, rows, path):
    wb = Workbook()
    ws = wb.active
    ws.title = "Results"

    ws.append(headers)

    green_fill = PatternFill(start_color='C6EFCE', end_color='C6EFCE', fill_type='solid')  # Light green
    red_fill = PatternFill(start_color='FFC7CE', end_color='FFC7CE', fill_type='solid')    # Light red

    for row_data in rows:
        ws.append(row_data)
        status = str(row_data[-1]).lower()
        fill = green_fill if "qualified" in status and "not" not in status else red_fill

        for cell in ws[ws.max_row]:
            cell.fill = fill

    wb.save(path)

def audio_transcription(file_path):
    transcriber = aai.Transcriber()
    try:
        transcript = transcriber.transcribe(file_path)
        if transcript.error:
            return f"Transcription error: {transcript.error}"
        else:
            print("Transcribed Text: ", transcript)
        return transcript.text
    except Exception as e:
        return f"Transcription failed: {str(e)}"

def build_prompt(transcript, questions):
    numbered_questions = '\n'.join([f"{i+1}. {q}" for i, q in enumerate(questions)])
    return f"""
You are an AI assistant. Below is a transcript of a call. Based on this, please extract the answers to the following questions, If any question is not answered, say Customer didnot answer. And if this question was not present in the script, then reply that Agent didn't ask 

Questions:
{numbered_questions}

Transcript:
{transcript}

Please return answers in this format:
Q1: [answer]
Q2: [answer]
Q3: [answer]
Q4: [answer]
Q5: [answer]
"""

def resource_generation_gemini(prompt_text):
    apikey = "AIzaSyC3akknXGbeuvAg_kBRkGsi586RuXeHkHo"  # Replace with your actual Gemini API key
    url = f"https://generativelanguage.googleapis.com/v1beta/models/gemini-2.0-flash:generateContent?key={apikey}"
    headers = {"Content-Type": "application/json"}
    data = {
        "contents": [
            {
                "parts": [{"text": prompt_text}]
            }
        ]
    }

    try:
        response = requests.post(url, headers=headers, json=data)
        response.raise_for_status()
        res_json = response.json()
        candidates = res_json.get("candidates", [])
        if candidates and "content" in candidates[0]:
            parts = candidates[0]["content"].get("parts", [])
            if parts:
                raw_text = parts[0].get("text", "")
                return parse_gemini_response(raw_text)
        return {}
    except requests.exceptions.RequestException as e:
        print(f"Gemini API request failed: {e}")
        return {}

def parse_gemini_response(text):
    import re
    answers = {}
    lines = text.splitlines()

    for line in lines:
        match = re.match(r"Q(\d)[:.\)]\s*(.+)", line.strip(), re.IGNORECASE)
        if match:
            q_index = int(match.group(1)) - 1
            answer = match.group(2).strip()
            questions_map = {
                0: "what is the year make and model of your vehicle?",
                1: "how many miles does your vehicle have",
                2: "can you spell your first and last name, please",
                3: "in witch state do you currently reside",
                4: "may i have your email address, please"
            }
            answers[questions_map.get(q_index, f"q{q_index+1}")] = answer

    return answers

def is_qualified_gemini(answers):
    with open("script1.txt", "r") as file:
        scriptdata = file.read()
    
    instructions = "You have a script with questions and there required answers and you also have the answers answered by a customer, you have to match the answers. If all answers are according to the script and none of them is missing, you have to reply Qualified(keep in mind that all questions must have answers, if agent didn't asked and customer didn't reply comes, then customer is not qulaified and email is optional, if email is not present, you can proceed further eith qualified status), else you have to reply Not Qualified follwed by a small desc of reason in 4-5 words only." 
    prompt_qualify = instructions + "Script: " + 'Question 1: What is the year make and model of your vehicle? Answer: (Year 2007 to 2024) Not cover -Luxury vehicles like Porsche 911 GT3, Lamborghini all models, Rolls-Royce all model Bentley models, Electric or lease vehicle, tesla, sports car, all expensive vehicle Question 2: How many miles does your vehicle have Answer: Minimum 1100 / Maximum 200000 Question 3: Can you spell your first and last name, please Question 4: In witch state do you currently reside Answer: Question: May I have your email address, please States wo don’t cover are: 1.Missouri 2.Florida 3. Washington 4. New Mexico 5. California Cars we don’t cover are Electric cars in the U.S. Sports Cars in the US Tesla Model 3 2024 Nissan Z Chevrolet Bolt BMW Z4 Hyundai Ioniq 5 Mazda Miata Volkswagen ID4 2024 Chevrolet Camaro Chevrolet Bolt EUV Porsche 911 Hyundai Ioniq 6 Porsche Boxster Kia EV6 Porsche Cayman Polestar 2 Subaru BRZ BMW i4 Toyota GR86 Toyota Supra Hyundai Kona Electric Fiat124 Spider Nissan Leaf 2012 Audi R8 Porsche Taycan Audi tt BMW M2 Tesla Model Y BMW i5 Ford Mustang Mach 4x Gt Performance GMC Hummer EV Lucid Air Mercedes Benz Eqs 580 Front Rivian R1T Fisker Ocean MINI Cooper SE Qualified CAR YEARS: 2007,2008,2009,2010,2011,2012,2013,2014,2015,2016,2017,2018,2019,2020,2021,2022,2023,2024 Qualify Vehicles Make: ACURA AUDI BMW BUICK CADILLAC CHEVEROLET CHRYSLER DODGE FIAT FORD GENESIS GMC HONDA HUMMER HYUNDAI INFINITI ISUZU JAGUAR JEEP KIA LAND ROVER LEXUS LINCOLN MAZDA MERCEDES MINI MITSUBISHI MURCURY NISSAN PONTIAC RAM SATRURN SUBARU TOYOTA VAXHAULL VOLKSAWAGON VOLVO COLOUR: BLACK, BROWN, GREEN, BLUE, LIGHT BLUE, HAZEL, AMBER, GRAY States We cover: ALASKA ALABAMA ARKANSAS AMERICAN SAMOA ARIZONA COLORADO CONNECTICUT DISTRICT OF COLUMBIA DELAWARE GEORGIA HAWAII LOWA LDAHO ILLINOIS INDIANA KANSAS KENTUCKY LOUISIANA MASSACHUSETTS MARYLAND MAINE MICHIGAN MIISSISSIPPI MONTANA NORTHH CAROLINA NORTH DAKOTA NEBRASKA NEW HAMPSHIRE NEW JERSEY NEVADA NEW YORK OHIO OKLAHOMA OREGON PENNSYLVANIA PUERTO RICO RHODE ISLAND SOUTH CAROLINA SOUTH DAKOTA TENNESSEE TEXAS UTAH VIRGINIA VIRGIN ISLANDS VERMONT WISCONSIN WEST VIRGINIA WYOMING' + "Call with Q/A: " + json.dumps(answers, indent=2)
    
    apikey = "AIzaSyC3akknXGbeuvAg_kBRkGsi586RuXeHkHo"  # Replace with your actual Gemini API key
    url = f"https://generativelanguage.googleapis.com/v1beta/models/gemini-2.0-flash:generateContent?key={apikey}"
    headers = {"Content-Type": "application/json"}
    data = {
        "contents": [
            {
                "parts": [{"text": prompt_qualify}]
            }
        ]
    }

    try:
        response = requests.post(url, headers=headers, json=data)
        response.raise_for_status()
        res_json = response.json()
        candidates = res_json.get("candidates", [])
        if candidates and "content" in candidates[0]:
            parts = candidates[0]["content"].get("parts", [])
            if parts:
                raw_text = parts[0].get("text", "")
                return raw_text
        return {}
    except requests.exceptions.RequestException as e:
        print(f"Gemini API request failed: {e}")
        return {}

def is_qualified(answers): 
    qualify_states = {"texas", "ohio", "georgia", "arizona", "nevada", "tennessee", "north carolina", "illinois", "michigan", "pennsylvania"}  # Example qualifying states
    disqualify_states = {"missouri", "florida", "washington", "new mexico", "california"}
    
    qualify_makes = {
        "toyota", "honda", "ford", "chevrolet", "nissan", "hyundai", "kia", "jeep", "dodge", "gmc"
    }
    disqualify_makes = {
        "porsche", "lamborghini", "rolls-royce", "bentley", "tesla", "fisker", "lucid", "rivian"
    }
    
    disqualify_keywords = ["electric", "lease", "sports", "expensive"]
    valid_years = {str(y) for y in range(2007, 2025)}
    min_miles = 1100
    max_miles = 200000

    q1 = answers.get("what is the year make and model of your vehicle?", "").lower()
    q2 = answers.get("how many miles does your vehicle have", "").replace(",", "").lower()
    q4 = answers.get("in witch state do you currently reside", "").lower()

    # Check for valid state
    if not any(state in q4 for state in qualify_states):
        return False

    # Disqualify if state explicitly disqualified (safety check)
    if any(state in q4 for state in disqualify_states):
        return False

    # Check vehicle year
    year_found = any(y in q1 for y in valid_years)

    # Check vehicle make qualification
    make_qualified = any(make in q1 for make in qualify_makes)
    make_disqualified = any(bad_make in q1 for bad_make in disqualify_makes)

    # Disqualify on keywords
    keyword_disqualified = any(word in q1 for word in disqualify_keywords)

    if not year_found or not make_qualified or make_disqualified or keyword_disqualified:
        return False

    # Check mileage
    try:
        miles = int(''.join(filter(str.isdigit, q2)))
        if miles < min_miles or miles > max_miles:
            return False
    except ValueError:
        return False

    return True

if __name__ == "__main__":
    app.run(debug=True)