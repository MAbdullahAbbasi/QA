from flask import *
import os
import tempfile
from openpyxl import Workbook
from openpyxl.styles import PatternFill
import uuid
import assemblyai as aai
import subprocess
from io import BytesIO
from werkzeug.utils import secure_filename
from concurrent.futures import ThreadPoolExecutor
import json
import re

# Setup Flask app
template_dir = os.path.abspath(os.path.join(os.path.dirname(__file__), 'templates'))
app = Flask(__name__, template_folder=template_dir, static_folder='../static')

EXPORT_FOLDER = os.path.join(os.getcwd(), 'exports')
os.makedirs(EXPORT_FOLDER, exist_ok=True)

UPLOAD_FOLDER = os.path.join(os.getcwd(), 'uploads')
os.makedirs(UPLOAD_FOLDER, exist_ok=True)

# API Keys
aai.settings.api_key = "f2b25c07671e4d2cac25479469debcab"

@app.route("/")
def home():
    return render_template("main.html")

def audio_transcription(file_path):
    transcriber = aai.Transcriber()
    try:
        transcript = transcriber.transcribe(file_path)
        if transcript.error:
            return f"Transcription error: {transcript.error}"
        return transcript.text
    except Exception as e:
        return f"Transcription failed: {str(e)}"

def build_prompt(transcript, questions):
    numbered_questions = '\n'.join([f"{i+1}. {q}" for i, q in enumerate(questions)])
    return f"""
You are an AI assistant. Below is a transcript of a call. Based on this, please extract the answers to the following questions, If any question is not answered, say Customer didnot answer. And if this question was not present in the script, then reply that Agent didn't ask 

Questions:
{numbered_questions}

Transcript:
{transcript}

Please return answers in this format:
Q1: [answer]
Q2: [answer]
Q3: [answer]
Q4: [answer]
Q5: [answer]
"""

def call_ollama_model(prompt):
    try:
        result = subprocess.run([
            "ollama", "run", "mistral"
        ], input=prompt.encode(), capture_output=True, timeout=120)
        output = result.stdout.decode().strip()
        return output
    except Exception as e:
        print(f"Ollama call failed: {e}")
        return ""

def resource_generation_ollama(prompt_text):
    raw_text = call_ollama_model(prompt_text)
    return parse_gemini_response(raw_text)

def parse_gemini_response(text):
    answers = {}
    lines = text.splitlines()
    for line in lines:
        match = re.match(r"Q(\d)[:.\)]\s*(.+)", line.strip(), re.IGNORECASE)
        if match:
            q_index = int(match.group(1)) - 1
            answer = match.group(2).strip()
            questions_map = {
                0: "what is the year make and model of your vehicle?",
                1: "how many miles does your vehicle have",
                2: "can you spell your first and last name, please",
                3: "in witch state do you currently reside",
                4: "may i have your email address, please"
            }
            answers[questions_map.get(q_index, f"q{q_index+1}")] = answer
    return answers

def is_qualified_ollama(answers):
    instructions = """
    You are given:
    1. A script containing required questions and expected answers.
    2. A set of customer responses.

    Your task:
    - Match each customer answer to the corresponding question in the script.
    - If all questions have been answered as expected, respond: Qualified.
    - If any question is missing or not properly answered, respond: Not Qualified followed by a short reason (4-5 words only).

    Important notes:
    - If the agent skipped a question and the customer did not answer it, the result is Not Qualified.
    - Email is optional.
    """
    prompt = instructions + "\nCall with Q/A: " + json.dumps(answers, indent=2)
    return call_ollama_model(prompt)

def save_excel(headers, rows, path):
    wb = Workbook()
    ws = wb.active
    ws.title = "Results"
    ws.append(headers)

    green_fill = PatternFill(start_color='C6EFCE', end_color='C6EFCE', fill_type='solid')
    red_fill = PatternFill(start_color='FFC7CE', end_color='FFC7CE', fill_type='solid')

    for row_data in rows:
        ws.append(row_data)
        status = str(row_data[-1]).lower()
        fill = green_fill if "qualified" in status and "not" not in status else red_fill
        for cell in ws[ws.max_row]:
            cell.fill = fill

    wb.save(path)

def process_audio_file(audio_file, questions):
    filename = secure_filename(audio_file.filename)
    with tempfile.NamedTemporaryFile(delete=False, suffix='.wav') as temp_audio:
        audio_file.save(temp_audio.name)
        transcribed_text = audio_transcription(temp_audio.name)
    os.remove(temp_audio.name)

    prompt = build_prompt(transcribed_text, questions)
    filtered_answers = resource_generation_ollama(prompt)
    qualification = is_qualified_ollama(filtered_answers)
    answers = [filtered_answers.get(q.lower(), '') for q in questions]
    return [filename] + answers + [qualification]

@app.route('/process_folder', methods=['POST'])
def process_folder():
    uploaded_files = request.files.getlist('audios')
    if not uploaded_files:
        return jsonify({"error": "No files uploaded"}), 400

    questions = [
        "What is the year make and model of your vehicle?",
        "How many miles does your vehicle have",
        "Can you spell your first and last name, please",
        "In witch state do you currently reside",
        "May I have your email address, please"
    ]

    headers = [
        "File Name", "Vehicle Info", "Mileage", "Full Name", "State", "Email", "Qualification Status"
    ]

    with ThreadPoolExecutor(max_workers=5) as executor:
        results = list(executor.map(lambda f: process_audio_file(f, questions), uploaded_files))

    filename = f'results_{uuid.uuid4().hex}.xlsx'
    excel_path = os.path.join(EXPORT_FOLDER, filename)
    save_excel(headers, results, excel_path)

    return jsonify({
        "headers": headers,
        "rows": results,
        "excel_download_url": url_for('download_excel', filename=filename)
    })

@app.route('/process_files', methods=['POST'])
def process_files():
    return process_folder()

@app.route('/download_excel/<filename>')
def download_excel(filename):
    file_path = os.path.join(EXPORT_FOLDER, filename)
    if not os.path.exists(file_path):
        return jsonify({"error": "File not found"}), 404
    return send_file(file_path, mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
                     as_attachment=True, download_name='results.xlsx')

if __name__ == "__main__":
    app.run(debug=True)